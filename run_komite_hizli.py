# -*- coding: utf-8 -*-
"""
HIZLI Komite/Komisyon Raporları İşleme
OCR atlanır → PDF'lerden doğrudan metin çıkarılır
Metin yoksa dosya adı bilgisinden kayıt oluşturulur
Tahmini süre: ~10-15 dakika (6800+ dosya)
"""
import os
import sys
import re
import time
import sqlite3
from datetime import datetime
from config import KOMITE_DIR
# Import get_db from app.py for multi-DB support
from app import get_db
from reference_matcher import ReferenceMatcher


def format_sure(sn):
    if sn < 60:
        return f"{int(sn)} saniye"
    elif sn < 3600:
        return f"{int(sn//60)} dk {int(sn%60)} sn"
    else:
        return f"{int(sn//3600)} saat {int((sn%3600)//60)} dk"


def parse_tarih(text):
    if not text:
        return None
    for pat in [r'(\d{2})[./](\d{2})[./](\d{4})', r'(\d{4})[./](\d{2})[./](\d{2})']:
        m = re.search(pat, str(text))
        if m:
            g = m.groups()
            try:
                if len(g[0]) == 4:
                    return datetime.strptime(f"{g[0]}-{g[1]}-{g[2]}", "%Y-%m-%d").date()
                else:
                    return datetime.strptime(f"{g[0]}/{g[1]}/{g[2]}", "%d/%m/%Y").date()
            except ValueError:
                continue
    return None


def parse_saat(text):
    if not text:
        return None
    m = re.search(r'(\d{1,2})[:.](\d{2})', str(text))
    if m:
        try:
            return datetime.strptime(f"{m.group(1)}:{m.group(2)}", "%H:%M").time()
        except ValueError:
            pass
    return None


def detect_uygunluk(text):
    if not text:
        return None
    t = text.lower().strip()
    if 'kısmen' in t:
        return 'Kısmen Karşılanıyor'
    elif 'karşılanmıyor' in t:
        return 'Karşılanmıyor'
    elif 'karşılanıyor' in t:
        return 'Karşılanıyor'
    return text.strip()[:30] if text.strip() else None


def detect_son_durum(text):
    if not text:
        return None
    t = text.lower().strip()
    if 'tamamlandı' in t and 'tamamlanmadı' not in t:
        return 'Tamamlandı'
    elif 'tamamlanmadı' in t:
        return 'Tamamlanmadı'
    elif 'devam' in t:
        return 'Devam Ediyor'
    return text.strip()[:30] if text.strip() else None


def hizli_docx_parse(filepath, file_info, matcher):
    """DOCX dosyasını hızlı parse et"""
    from docx import Document
    try:
        doc = Document(filepath)
    except Exception:
        return None, [], None

    rapor = {
        'kurum_kodu': file_info.get('kurum_kodu'),
        'il': file_info.get('il', ''),
        'ilce': file_info.get('ilce', ''),
        'hastane_adi': file_info.get('hastane', '') or '',
        'rapor_tipi': file_info.get('rapor_tipi', 'Komisyon'),
        'degerlendirme_tarihi': None,
        'degerlendirme_saati': None,
        'ekip_uyeleri': '',
        'kaynak_dosya': os.path.basename(filepath),
        'dosya_formati': 'docx'
    }
    standartlar = []
    karar = None

    for ti, table in enumerate(doc.tables):
        rows = table.rows
        if not rows:
            continue

        if ti == 0:
            for row in rows:
                cells = [cell.text.strip() for cell in row.cells]
                if len(cells) >= 2:
                    label = cells[0].lower()
                    value = cells[1]
                    if 'tesis' in label and 'adı' in label:
                        if value:
                            hospital = matcher.match_by_name(value, il=rapor['il'])
                            if hospital:
                                rapor['hastane_adi'] = hospital['kurum_adi']
                                rapor['kurum_kodu'] = hospital['kurum_kodu']
                            else:
                                rapor['hastane_adi'] = value
                    elif 'kurum kodu' in label:
                        try:
                            rapor['kurum_kodu'] = int(value)
                        except (ValueError, TypeError):
                            pass
                    elif 'tarih' in label:
                        rapor['degerlendirme_tarihi'] = parse_tarih(value)
                    elif 'saat' in label:
                        rapor['degerlendirme_saati'] = parse_saat(value)
                    elif 'üye' in label or 'ekip' in label:
                        rapor['ekip_uyeleri'] = value
        else:
            try:
                header_cells = [cell.text.strip().lower() for cell in rows[0].cells]
            except:
                continue

            is_standart = any('standart' in h or 'uygunluk' in h or 'karşılan' in h for h in header_cells)

            if is_standart:
                for ri in range(1, len(rows)):
                    try:
                        cells = [cell.text.strip() for cell in rows[ri].cells]
                    except:
                        continue
                    if len(cells) < 2:
                        continue

                    s = {
                        'standart_no': '', 'standart_adi': '', 'degerlendirme_olcutu': '',
                        'uygunluk_durumu': None, 'eksikler': '', 'sorumlu': '',
                        'planlanan_baslangic_tarihi': None, 'planlanan_bitis_tarihi': None,
                        'son_durum': None, 'aciklama': ''
                    }
                    for ci, ct in enumerate(cells):
                        h = header_cells[ci % len(header_cells)] if ci < len(header_cells) else ''
                        if 'standart' in h and 'no' in h:
                            s['standart_no'] = ct
                        elif 'standart' in h:
                            s['standart_adi'] = ct
                        elif 'ölçüt' in h or 'kriter' in h:
                            s['degerlendirme_olcutu'] = ct
                        elif 'uygunluk' in h or 'karşılan' in h:
                            s['uygunluk_durumu'] = detect_uygunluk(ct)
                        elif 'eksik' in h:
                            s['eksikler'] = ct
                        elif 'sorumlu' in h:
                            s['sorumlu'] = ct
                        elif 'son durum' in h or 'durum' in h:
                            s['son_durum'] = detect_son_durum(ct)
                        elif 'tarih' in h:
                            s['planlanan_baslangic_tarihi'] = parse_tarih(ct)

                    if s['standart_no'] or s['standart_adi'] or s['uygunluk_durumu']:
                        standartlar.append(s)
            else:
                for row in rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    combined = ' '.join(cells).lower()
                    if 'iyileştirme' in combined or 'komisyon kararı' in combined:
                        if karar is None:
                            karar = {'iyilestirme_alanlari': '', 'komisyon_karari': '', 'muafiyetler': ''}
                        for ct in cells:
                            cl = ct.lower()
                            if 'iyileştirme' in cl and len(ct) > 30:
                                karar['iyilestirme_alanlari'] = ct
                            elif 'komisyon kararı' in cl and len(ct) > 20:
                                karar['komisyon_karari'] = ct
                            elif 'muaf' in cl:
                                karar['muafiyetler'] = ct
                            elif len(ct) > 50 and not karar['komisyon_karari']:
                                karar['komisyon_karari'] = ct

    if rapor['kurum_kodu'] and not rapor['hastane_adi']:
        hospital = matcher.match_by_code(rapor['kurum_kodu'])
        if hospital:
            rapor['hastane_adi'] = hospital['kurum_adi']

    return rapor, standartlar, karar


def hizli_pdf_parse(filepath, file_info, matcher):
    """PDF'den sadece metin çıkar (OCR YOK - çok hızlı)"""
    import fitz
    try:
        doc = fitz.open(filepath)
    except Exception:
        return None, [], None

    rapor = {
        'kurum_kodu': file_info.get('kurum_kodu'),
        'il': file_info.get('il', ''),
        'ilce': file_info.get('ilce', ''),
        'hastane_adi': file_info.get('hastane', '') or '',
        'rapor_tipi': file_info.get('rapor_tipi', 'Komite'),
        'degerlendirme_tarihi': None,
        'degerlendirme_saati': None,
        'ekip_uyeleri': '',
        'kaynak_dosya': os.path.basename(filepath),
        'dosya_formati': 'pdf'
    }
    standartlar = []

    full_text = ''
    for page in doc:
        full_text += page.get_text() + '\n'
    doc.close()

    if full_text.strip():
        rapor['degerlendirme_tarihi'] = parse_tarih(full_text)
        rapor['degerlendirme_saati'] = parse_saat(full_text)

        # Üye tespiti
        uye_list = re.findall(r'(?:üye|ekip|katılımcı|başkan)[\s:]*([^\n]+)', full_text, re.IGNORECASE)
        if uye_list:
            rapor['ekip_uyeleri'] = '\n'.join(set(uye_list))

        # Standart tespiti
        blocks = re.split(r'(?=Standart\s*\d)', full_text)
        for block in blocks:
            m = re.match(r'Standart\s*(\d+[.\d]*)', block)
            if m:
                s = {
                    'standart_no': m.group(1), 'standart_adi': block[:200].strip(),
                    'degerlendirme_olcutu': '', 'uygunluk_durumu': detect_uygunluk(block),
                    'eksikler': '', 'sorumlu': '',
                    'planlanan_baslangic_tarihi': None, 'planlanan_bitis_tarihi': None,
                    'son_durum': detect_son_durum(block), 'aciklama': ''
                }
                standartlar.append(s)

    # Hastane adı eşleştirme
    if rapor['kurum_kodu'] and not rapor['hastane_adi']:
        hospital = matcher.match_by_code(rapor['kurum_kodu'])
        if hospital:
            rapor['hastane_adi'] = hospital['kurum_adi']
            rapor['ilce'] = hospital['ilce']

    if not rapor['hastane_adi']:
        rapor['hastane_adi'] = f"{rapor['il']} - {rapor['ilce']} Sağlık Tesisi"

    return rapor, standartlar, None


def kaydet(rapor, standartlar, karar, cur, conn):
    """Raporu veritabanına kaydet"""
    kurum_kodu = rapor['kurum_kodu']
    from config import USE_SQLITE
    placeholder = '?' if USE_SQLITE else '%s'
    if kurum_kodu is not None:
        cur.execute(f"SELECT 1 FROM referans_hastaneler WHERE kurum_kodu = {placeholder}", (kurum_kodu,))
        if not cur.fetchone():
            kurum_kodu = None

    from config import USE_SQLITE
    placeholders = '?, ?, ?, ?, ?, ?, ?, ?, ?, ?' if USE_SQLITE else '%s, %s, %s, %s, %s, %s, %s, %s, %s, %s'
    
    cur.execute(f"""
        INSERT INTO komite_raporlari
        (kurum_kodu, il, ilce, hastane_adi, rapor_tipi, degerlendirme_tarihi,
         degerlendirme_saati, ekip_uyeleri, kaynak_dosya, dosya_formati)
        VALUES ({placeholders})
    """, (
        kurum_kodu, rapor['il'], rapor['ilce'], rapor['hastane_adi'],
        rapor['rapor_tipi'], rapor['degerlendirme_tarihi'],
        rapor['degerlendirme_saati'], rapor['ekip_uyeleri'],
        rapor['kaynak_dosya'], rapor['dosya_formati']
    ))
    
    if USE_SQLITE:
        rapor_id = cur.lastrowid
    else:
        rapor_id = cur.fetchone()[0]

    from config import USE_SQLITE
    s_placeholders = '?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?' if USE_SQLITE else '%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s'
    
    for s in standartlar:
        cur.execute(f"""
            INSERT INTO standart_degerlendirmeler
            (rapor_id, standart_no, standart_adi, degerlendirme_olcutu,
             uygunluk_durumu, eksikler, sorumlu,
             planlanan_baslangic_tarihi, planlanan_bitis_tarihi,
             son_durum, aciklama)
            VALUES ({s_placeholders})
        """, (
            rapor_id, s['standart_no'], s['standart_adi'],
            s['degerlendirme_olcutu'], s['uygunluk_durumu'],
            s['eksikler'], s['sorumlu'],
            s['planlanan_baslangic_tarihi'], s['planlanan_bitis_tarihi'],
            s['son_durum'], s['aciklama']
        ))

    if karar:
        from config import USE_SQLITE
        k_placeholders = '?, ?, ?, ?' if USE_SQLITE else '%s, %s, %s, %s'
        cur.execute(f"""
            INSERT INTO komisyon_kararlari
            (rapor_id, iyilestirme_alanlari, komisyon_karari, muafiyetler)
            VALUES ({k_placeholders})
        """, (rapor_id, karar['iyilestirme_alanlari'],
              karar['komisyon_karari'], karar['muafiyetler']))

    return rapor_id


def main():
    print()
    print("=" * 70)
    print("  ⚡ HIZLI KOMİTE/KOMİSYON RAPORLARI İŞLEME")
    print("  📌 OCR devre dışı → Doğrudan metin çıkarma ile hızlı işlem")
    print("=" * 70)
    print()

    files = sorted([
        f for f in os.listdir(KOMITE_DIR)
        if os.path.isfile(os.path.join(KOMITE_DIR, f))
    ])
    total = len(files)

    ext_counts = {}
    for f in files:
        ext = os.path.splitext(f)[1].lower()
        ext_counts[ext] = ext_counts.get(ext, 0) + 1

    print(f"  📁 Klasör : {KOMITE_DIR}")
    print(f"  📊 Toplam : {total} dosya")
    for ext, cnt in sorted(ext_counts.items(), key=lambda x: -x[1]):
        print(f"     {ext:6s} : {cnt} dosya")

    # Tahmini süre (OCR yok → PDF ~0.3sn, DOCX ~0.5sn, diğer ~0.05sn)
    est = ext_counts.get('.pdf', 0) * 0.3 + ext_counts.get('.docx', 0) * 0.5 + \
          (total - ext_counts.get('.pdf', 0) - ext_counts.get('.docx', 0)) * 0.05
    print(f"\n  ⏰ Tahmini süre: {format_sure(est)}")
    print()

    # Veritabanı temizle
    print("  🗑️  Mevcut veriler temizleniyor...")
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM komite_raporlari")
    old = cur.fetchone()[0]
    cur.execute("DELETE FROM komisyon_kararlari")
    cur.execute("DELETE FROM standart_degerlendirmeler")
    cur.execute("DELETE FROM komite_raporlari")
    conn.commit()
    print(f"  ✓ {old} eski kayıt temizlendi.\n")

    print("  🚀 İŞLEM BAŞLADI!")
    print("  " + "-" * 68)

    matcher = ReferenceMatcher()

    t0 = time.time()
    ok = 0
    err = 0
    tip = {'docx': 0, 'pdf': 0, 'xlsx': 0, 'doc': 0, 'img': 0, 'skip': 0}
    il_say = {}

    for i, fn in enumerate(files):
        fp = os.path.join(KOMITE_DIR, fn)
        ext = os.path.splitext(fn)[1].lower()
        fi = matcher.match_filename(fn)

        try:
            rapor = None
            standartlar = []
            karar = None

            if ext == '.docx':
                rapor, standartlar, karar = hizli_docx_parse(fp, fi, matcher)
                if rapor:
                    kaydet(rapor, standartlar, karar, cur, conn)
                    conn.commit()
                    tip['docx'] += 1

            elif ext == '.pdf':
                rapor, standartlar, karar = hizli_pdf_parse(fp, fi, matcher)
                if rapor:
                    kaydet(rapor, standartlar, karar, cur, conn)
                    conn.commit()
                    tip['pdf'] += 1

            elif ext in ('.xlsx', '.xls'):
                rapor = {
                    'kurum_kodu': fi.get('kurum_kodu'), 'il': fi.get('il', ''),
                    'ilce': fi.get('ilce', ''),
                    'hastane_adi': fi.get('hastane', '') or f"{fi.get('il', '')} Sağlık Tesisi",
                    'rapor_tipi': fi.get('rapor_tipi', 'Komisyon'),
                    'degerlendirme_tarihi': None, 'degerlendirme_saati': None,
                    'ekip_uyeleri': '', 'kaynak_dosya': fn, 'dosya_formati': 'xlsx'
                }
                kaydet(rapor, [], None, cur, conn)
                conn.commit()
                tip['xlsx'] += 1

            elif ext in ('.jpg', '.jpeg', '.png', '.bmp', '.tif', '.tiff', '.jfif'):
                rapor = {
                    'kurum_kodu': fi.get('kurum_kodu'), 'il': fi.get('il', ''),
                    'ilce': fi.get('ilce', ''),
                    'hastane_adi': fi.get('hastane', '') or f"{fi.get('il', '')} Sağlık Tesisi",
                    'rapor_tipi': fi.get('rapor_tipi', 'Komite'),
                    'degerlendirme_tarihi': None, 'degerlendirme_saati': None,
                    'ekip_uyeleri': '', 'kaynak_dosya': fn, 'dosya_formati': ext.replace('.', '')
                }
                kaydet(rapor, [], None, cur, conn)
                conn.commit()
                tip['img'] += 1

            elif ext == '.doc':
                rapor = {
                    'kurum_kodu': fi.get('kurum_kodu'), 'il': fi.get('il', ''),
                    'ilce': fi.get('ilce', ''),
                    'hastane_adi': fi.get('hastane', '') or f"{fi.get('il', '')} Sağlık Tesisi",
                    'rapor_tipi': fi.get('rapor_tipi', 'Komisyon'),
                    'degerlendirme_tarihi': None, 'degerlendirme_saati': None,
                    'ekip_uyeleri': '', 'kaynak_dosya': fn, 'dosya_formati': 'doc'
                }
                kaydet(rapor, [], None, cur, conn)
                conn.commit()
                tip['doc'] += 1

            else:
                tip['skip'] += 1
                continue

            if rapor:
                ok += 1
                r_il = rapor.get('il', '?') or '?'
                il_say[r_il] = il_say.get(r_il, 0) + 1

        except Exception as e:
            err += 1
            conn.rollback()
            if err <= 15:
                print(f"  ⚠️  HATA: {fn[:50]}... → {str(e)[:80]}")

        # Her 10 dosyada rapor
        n = i + 1
        if n % 10 == 0 or n == total:
            el = time.time() - t0
            hiz = n / el if el > 0 else 0
            kalan = (total - n) / hiz if hiz > 0 else 0
            pct = n / total * 100
            bar_w = 40
            done = int(bar_w * n / total)
            bar = '█' * done + '░' * (bar_w - done)

            il_now = fi.get('il', '?') or '?'
            print(f"\r  [{bar}] %{pct:5.1f}  |  {n}/{total}  |  ✅{ok} ❌{err}  |  ⏱️{format_sure(el)} kalan:{format_sure(kalan)}  |  📍{il_now}  |  🚀{hiz:.1f}/sn", end='', flush=True)

    cur.close()
    conn.close()

    toplam_sure = time.time() - t0
    print()
    print()
    print("  " + "=" * 68)
    print("  ✅ İŞLEM TAMAMLANDI!")
    print("  " + "=" * 68)
    print()
    print(f"  📊 SONUÇLAR:")
    print(f"     Toplam dosya      : {total}")
    print(f"     Başarılı          : {ok}")
    print(f"     Hata              : {err}")
    print(f"     Atlanan           : {tip['skip']}")
    print(f"     Toplam süre       : {format_sure(toplam_sure)}")
    print(f"     Ortalama hız      : {ok/toplam_sure:.1f} dosya/sn")
    print()
    print(f"  📋 DOSYA TÜRLERİ:")
    print(f"     PDF   : {tip['pdf']}")
    print(f"     DOCX  : {tip['docx']}")
    print(f"     XLSX  : {tip['xlsx']}")
    print(f"     DOC   : {tip['doc']}")
    print(f"     Görsel : {tip['img']}")
    print()
    print(f"  🗺️  İL DAĞILIMI ({len(il_say)} il):")
    for il, cnt in sorted(il_say.items(), key=lambda x: -x[1]):
        b = '▓' * min(cnt // 10, 30)
        print(f"     {il:20s} : {cnt:5d} {b}")
    print()
    print("  " + "=" * 68)
    sys.stdout.flush()


if __name__ == '__main__':
    main()
